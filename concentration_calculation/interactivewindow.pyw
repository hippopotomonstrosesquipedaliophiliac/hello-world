# interactivewindow.pyw
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import os
from pathlib import Path
from openpyxl import load_workbook
import pandas as pd
import sys
from matplotlib.backends.backend_tkagg import NavigationToolbar2Tk, FigureCanvasTkAgg

# Plotting inside Tk
import matplotlib.pyplot as plt

# --- make paths stable no matter where this file is launched from ---
REACTION_EXTENT_PREFIX = "__eq_x__::"  # keys in iteration_history for per-reaction x

BASE_DIR = Path(__file__).resolve().parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

# --- your domain logic modules ---
import system_reaction_parser as srp   # CompartmentSystem
import helpertesting as hd             # SpeciesMatrix + helpers
from testing_miscellaneous import (
    check_if_previous_data_frame_exist,
    edit_then_wait,
)

# ------------------------- App Shell -------------------------
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Concentration Pipeline — Big Window")
        self.geometry("1000x650")
        
        # Shared state across frames
        self.state = {
            "docx_path": tk.StringVar(value=""),
            "reactants_folder": tk.StringVar(
                value=str((BASE_DIR / "reactants' dataframe").resolve())
            ),
            "chosen_xlsx": tk.StringVar(value=""),
            "systems_ready": tk.BooleanVar(value=False),
            "summary": tk.StringVar(value=""),
        }

        # Domain objects (populated later)
        self.compartment_manager = None  # srp.CompartmentSystem
        self.previous_matches = []       # list[str] of prior xlsx names

        # Frame order (linear)
        self.frame_order = [
            "WelcomeFrame",
            "DocxFrame",
            "DataframeFrame",  # combined scan + inline create/edit
            "RunFrame",
        ]
        self.current_index = 0

        # Container
        self.container = ttk.Frame(self)
        self.container.pack(fill="both", expand=True)
        self.container.rowconfigure(0, weight=1)
        self.container.columnconfigure(0, weight=1)

        # Build frames
        self.frames = {}
        for name in self.frame_order:
            F = globals()[name]
            frame = F(parent=self.container, controller=self)
            self.frames[name] = frame
            frame.grid(row=0, column=0, sticky="nsew")

        # Menus + bottom nav
        self._build_menubar()
        self._build_bottom_nav()
        self.show_frame(self.frame_order[self.current_index])
        self.protocol("WM_DELETE_WINDOW", self._on_app_close)
    def _on_app_close(self):
        """Hard-close all child windows, Matplotlib figures, and the Tk mainloop."""
        # 1) Close any Matplotlib figures
        try:
            import matplotlib.pyplot as plt
            plt.close("all")
        except Exception:
            pass

        # 2) Release grabs and destroy all Toplevel children
        for w in list(self.winfo_children()):
            try:
                # Some modals might have a grab; release safely
                try:
                    self.grab_release()
                except Exception:
                    pass
                w.destroy()
            except Exception:
                pass

        # 3) Quit + destroy the root
        try:
            self.quit()
        except Exception:
            pass
        try:
            self.destroy()
        except Exception:
            pass

    # --- Navigation ---
    def show_frame(self, name: str):
        frame = self.frames[name]
        frame.tkraise()
        if hasattr(frame, "on_show"):
            frame.on_show()
        if name in self.frame_order:
            self.current_index = self.frame_order.index(name)
        self._update_nav_buttons()

    def next_frame(self):
        if self.current_index < len(self.frame_order) - 1:
            self.current_index += 1
            self.show_frame(self.frame_order[self.current_index])

    def prev_frame(self):
        if self.current_index > 0:
            self.current_index -= 1
            self.show_frame(self.frame_order[self.current_index])

    def _build_menubar(self):
        menubar = tk.Menu(self)
        pages_menu = tk.Menu(menubar, tearoff=0)
        for key in self.frame_order:
            label = key.replace("Frame", "")
            pages_menu.add_command(label=label, command=lambda n=key: self.show_frame(n))
        menubar.add_cascade(label="Pages", menu=pages_menu)
        self.config(menu=menubar)

    def _build_bottom_nav(self):
        bottom = ttk.Frame(self)
        bottom.pack(fill="x", side="bottom")
        # left status
        self.step_label = ttk.Label(bottom, text="Step 1 of {}".format(len(self.frame_order)))
        self.step_label.pack(side="left", padx=8, pady=6)
        # right buttons — Next first, then Back (for < > visual order)
        self.next_btn = ttk.Button(bottom, text="Next  ▶", command=self.next_frame)
        self.next_btn.pack(side="right", padx=6, pady=6)
        self.prev_btn = ttk.Button(bottom, text="◀  Back", command=self.prev_frame)
        self.prev_btn.pack(side="right", padx=(0, 6), pady=6)

    def _update_nav_buttons(self):
        at_head = self.current_index == 0
        at_tail = self.current_index == len(self.frame_order) - 1
        self.prev_btn.state(["disabled"] if at_head else ["!disabled"])
        self.next_btn.state(["disabled"] if at_tail else ["!disabled"])
        self.step_label.config(text=f"Step {self.current_index+1} of {len(self.frame_order)}")

    # -------------------- Pipeline steps --------------------
    def parse_docx(self):
        path = self.state["docx_path"].get().strip()
        if not path:
            messagebox.showerror("Missing DOCX", "Please choose a .docx first.")
            return False
        try:
            self.compartment_manager = srp.CompartmentSystem(path)
            self.state["systems_ready"].set(True)
            return True
        except Exception as e:
            messagebox.showerror("Parse error", str(e))
            return False

    def detect_previous_frames(self):
        folder = self.state["reactants_folder"].get().strip() or str(BASE_DIR)
        Path(folder).mkdir(parents=True, exist_ok=True)

        # Collect current species list across systems
        systems = self.compartment_manager.get_all_systems() if self.compartment_manager else {}
        current_species = []
        for sys_name, info in systems.items():
            df = info["speciesMatrix"].construct_concentration_data(system_name=sys_name)
            current_species.extend(list(df["species"]))
        current_species = sorted(set(current_species))

        # Scan folder for existing .xlsx files and compare species set
        matches = []
        prev = check_if_previous_data_frame_exist(folders_path=folder)
        for d in prev:
            for fname, species_list in d.items():
                if current_species == sorted(species_list):  # order-independent
                    matches.append(fname)
        self.previous_matches = matches
        return matches

    def assign_concentrations(self):
        xlsx = self.state["chosen_xlsx"].get().strip()
        if not xlsx:
            messagebox.showerror("Missing file", "No concentration workbook chosen.")
            return False
        # Read sheet names, then close (avoid Windows file lock)
        wb = load_workbook(xlsx, read_only=True)
        try:
            sheet_names = wb.sheetnames
        finally:
            wb.close()
        it = iter(sheet_names)
        for sys_name, info in self.compartment_manager.get_all_systems().items():
            try:
                sheet = next(it)
            except StopIteration:
                messagebox.showerror("Sheet mismatch", "More systems than sheets in workbook.")
                return False
            info["speciesMatrix"].appending_concentration_data(
                concentration_file_path=xlsx,
                system_sheet_name=sheet,
            )
        return True

    def run_compute_and_export(self):
        cm = self.compartment_manager
        # 1) Snapshot global concentrations
        _ = cm.snapshot_global_concentrations()
        # 2) Validate / assert reactions
        cm.validate_local_participation()
        cm.assert_local_participation()
        # 3) Advance all systems together to equilibrium-ish
        final_global = cm.solve_all_systems_iterative(max_iters=500, zero_tol=1e-5, window=20)

        # 4) Export per-system artifacts (unchanged)
        for sys_name, info in cm.get_all_systems().items():
            sm = info["speciesMatrix"]
            comp = info["compartment"]
            # templates / dataframes (your current exports)
            sm.export_species_template(output_file=f"system_{comp}_{sys_name}_template.xlsx")
            sm.get_species_dataframe(output_file=f"system_{comp}_{sys_name}_dataframe.xlsx")
        return True


    # --- history extraction helper for plotting ---
    def _get_iteration_history(self):
        """
        Collect iteration history across systems.

        Supports (per speciesMatrix):
          - get_iteration_history() -> {species: [(iter, conc), ...]}
          - iteration_history / history / convergence_log (same shape or [conc1, ...])

        Returns:
          { system_name: { "compartment": str, "series": {species: [(i, c), ...]}}}
        """
        if not self.compartment_manager:
            return {}

        history = {}
        for sys_name, info in self.compartment_manager.get_all_systems().items():
            sm = info["speciesMatrix"]
            series = None

            if hasattr(sm, "get_iteration_history") and callable(sm.get_iteration_history):
                try:
                    series = sm.get_iteration_history()
                except Exception:
                    series = None

            if series is None:
                for attr in ("iteration_history", "history", "convergence_log"):
                    cand = getattr(sm, attr, None)
                    if cand:
                        series = cand
                        break

            norm = {}
            if isinstance(series, dict):
                for sp, seq in series.items():
                    if seq and not isinstance(seq[0], (tuple, list)):
                        seq = list(enumerate(seq, start=1))  # [c1,c2] -> [(1,c1),(2,c2)]
                    norm[str(sp)] = [(int(i), float(c)) for i, c in seq]
            if norm:
                history[sys_name] = {"compartment": info.get("compartment", ""), "series": norm}
        return history

    def show_convergence_plots(self):
        history = self._get_iteration_history()
        if not history:
            messagebox.showinfo(
                "No iteration data",
                "No iteration history was found. Ensure your solver records per-iteration "
                "concentrations (e.g., SpeciesMatrix.iteration_history or get_iteration_history())."
            )
            return
        ConvergencePlotWindow(self, history)

# ------------------------- Convergence Plot Window -------------------------
# at top of file you already have:
# from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
# import matplotlib.pyplot as plt


class ConvergencePlotWindow(tk.Toplevel):
    def __init__(self, controller: App, history: dict):
        super().__init__(controller)
        self.title("Convergence — Concentration vs Iteration")
        self.controller = controller
        self.history = history
        self._mpl_cids = []   # mpl connection ids to clean up
        self._pan_state = None  # (x0,y0,xlim0,ylim0) in data coords while dragging

        # size to parent
        controller.update_idletasks()
        w = controller.winfo_width() or 1000
        h = controller.winfo_height() or 650
        x = controller.winfo_rootx()
        y = controller.winfo_rooty()
        self.geometry(f"{w}x{h}+{x}+{y}")

        # Controls
        bar = ttk.Frame(self); bar.pack(fill="x", padx=12, pady=8)
        ttk.Label(bar, text="System:").pack(side="left")
        self.var_system = tk.StringVar()
        systems = list(history.keys())
        sys_combo = ttk.Combobox(bar, textvariable=self.var_system, values=systems, state="readonly", width=30)
        sys_combo.pack(side="left", padx=6)

        ttk.Label(bar, text="Series (multi-select):").pack(side="left", padx=(12, 0))

# display name -> real key (species name or __eq_x__::label)
        self._display_to_key = {}

        # Multi-select listbox + scrollbar
        self.list_species = tk.Listbox(bar, selectmode="extended", exportselection=False, height=8, width=36)
        self.list_species.pack(side="left", padx=6)
        sb = tk.Scrollbar(bar, orient="vertical", command=self.list_species.yview)
        sb.pack(side="left", fill="y")
        self.list_species.configure(yscrollcommand=sb.set)

        ttk.Button(bar, text="Plot selected", command=self._plot_selected).pack(side="left", padx=8)
        ttk.Button(bar, text="Plot all", command=self._plot_all).pack(side="left")



        # Manual zoom buttons
        ttk.Button(bar, text="Zoom In", command=lambda: self._zoom_button(1.2)).pack(side="left", padx=(18, 4))
        ttk.Button(bar, text="Zoom Out", command=lambda: self._zoom_button(1/1.2)).pack(side="left", padx=4)

        ttk.Button(bar, text="Close", command=self._on_close).pack(side="right")

        # Figure+Canvas+Toolbar
        self.fig, self.ax = plt.subplots()
        self.canvas = FigureCanvasTkAgg(self.fig, master=self)
        self.toolbar = NavigationToolbar2Tk(self.canvas, self)
        self.toolbar.update()
        self.canvas.get_tk_widget().pack(fill="both", expand=True)

        # Axes styling (grid, minor grid, ticks)
        self._style_axes()

        # Populate initial combo values
        if systems:
            self.var_system.set(systems[0])
            self._refresh_species()
        self.var_system.trace_add("write", lambda *a: self._refresh_species())

        # Mouse wheel zoom + drag-to-pan events
        self._mpl_cids.append(self.canvas.mpl_connect('scroll_event', self._on_scroll))
        self._mpl_cids.append(self.canvas.mpl_connect('button_press_event', self._on_press))
        self._mpl_cids.append(self.canvas.mpl_connect('motion_notify_event', self._on_motion))
        self._mpl_cids.append(self.canvas.mpl_connect('button_release_event', self._on_release))

        # Proper cleanup
        self.protocol("WM_DELETE_WINDOW", self._on_close)

    # ---------- axes cosmetics ----------
    def _style_axes(self):
        ax = self.ax
        ax.grid(True, which='major', alpha=0.35)
        ax.grid(True, which='minor', alpha=0.15)
        # integer ticker for iterations
        try:
            from matplotlib.ticker import MaxNLocator, AutoMinorLocator
            ax.xaxis.set_major_locator(MaxNLocator(integer=True, prune=None))
            ax.xaxis.set_minor_locator(AutoMinorLocator(2))
            ax.yaxis.set_minor_locator(AutoMinorLocator(2))
        except Exception:
            pass
        # slightly stronger spines
        for spine in ax.spines.values():
            spine.set_linewidth(1.0)
        ax.set_xlabel("Iteration")
        ax.set_ylabel("Concentration")

    # ---------- plotting ----------
    def _refresh_species(self):
        sys_name = self.var_system.get()
        series = self.history.get(sys_name, {}).get("series", {}) or {}

        # Build display labels, stripping the reaction prefix for pretty names
        name_map = {}
        for key in sorted(series.keys()):
            if key.startswith(REACTION_EXTENT_PREFIX):
                pretty = "[rxn] " + key.replace(REACTION_EXTENT_PREFIX, "").strip()
            else:
                pretty = key
            name_map[pretty] = key

        self._display_to_key = name_map

        # Fill the listbox
        self.list_species.delete(0, tk.END)
        for display in sorted(name_map.keys()):
            self.list_species.insert(tk.END, display)

    def _plot_selected(self):
        sys_name = self.var_system.get()
        if not sys_name:
            return
        series = self.history.get(sys_name, {}).get("series", {}) or {}

        sel_indices = self.list_species.curselection()
        if not sel_indices:
            return

        displays = [self.list_species.get(i) for i in sel_indices]
        keys = [self._display_to_key.get(d, d) for d in displays]

        self._clear_ax()
        saw_reaction = False
        for key in keys:
            data = series.get(key, [])
            if not data:
                continue
            xs, ys = zip(*data)
            if key.startswith(REACTION_EXTENT_PREFIX):
                saw_reaction = True
                label = "[rxn] " + key.replace(REACTION_EXTENT_PREFIX, "").strip()
            else:
                label = key
            self.ax.plot(xs, ys, marker="o", label=label)

        self.ax.set_xlabel("Iteration")
        self.ax.set_ylabel("Concentration / extent x" if saw_reaction else "Concentration")
        self.ax.legend(loc="best")
        self.ax.set_title(f"{sys_name} — selected")
        self.ax.relim(); self.ax.autoscale_view()
        self.canvas.draw_idle()

    def _clear_ax(self):
        self.ax.clear()
        self._style_axes()

    def _plot(self):
        sys_name = self.var_system.get()
        display = self.var_species.get()
        if not sys_name or not display:
            return

        # Map display label back to real key used in the series dict
        key = self._display_to_key.get(display, display)
        data = self.history[sys_name]["series"].get(key, [])
        if not data:
            return

        self._clear_ax()
        xs, ys = zip(*data)
        self.ax.plot(xs, ys, marker="o")

        # Y-axis label: concentration for species, extent for reactions
        if key.startswith(REACTION_EXTENT_PREFIX):
            self.ax.set_ylabel("Reaction extent x")
            legend_label = display.replace("[rxn] ", "")
        else:
            self.ax.set_ylabel("Concentration")
            legend_label = display

        self.ax.set_title(f"{sys_name} — {legend_label}")
        self.ax.relim(); self.ax.autoscale_view()
        self.canvas.draw_idle()


    
    def _plot_all(self):
        sys_name = self.var_system.get()
        if not sys_name:
            return
        series = self.history[sys_name]["series"]
        if not series:
            return

        self._clear_ax()
        saw_reaction = False
        for real_key, data in series.items():
            if not data:
                continue
            xs, ys = zip(*data)
            # Make a pretty legend label
            if real_key.startswith(REACTION_EXTENT_PREFIX):
                saw_reaction = True
                label = real_key.replace(REACTION_EXTENT_PREFIX, "").strip()
                label = f"[rxn] {label}" if label else "[rxn]"
            else:
                label = real_key
            self.ax.plot(xs, ys, marker="o", label=label)

        # If any reaction series plotted, use a neutral ylabel
        self.ax.set_ylabel("Concentration / extent x" if saw_reaction else "Concentration")
        self.ax.legend(loc="best")
        self.ax.set_title(f"{sys_name} — all series")
        self.ax.relim(); self.ax.autoscale_view()
        self.canvas.draw_idle()


    # ---------- zoom / pan ----------
    def _zoom_button(self, scale):
        if not self.fig.axes:
            return
        ax = self.ax
        x0, x1 = ax.get_xlim()
        y0, y1 = ax.get_ylim()
        cx = 0.5 * (x0 + x1)
        cy = 0.5 * (y0 + y1)
        self._zoom_at(ax, cx, cy, scale)
        self.canvas.draw_idle()

    def _reset_view(self):
        self.ax.relim(); self.ax.autoscale_view()
        self.canvas.draw_idle()

    def _on_scroll(self, event):
        if event.inaxes != self.ax:
            return
        scale = 1.2 if getattr(event, "step", 1) > 0 else (1/1.2)
        self._zoom_at(event.inaxes, event.xdata, event.ydata, scale)
        self.canvas.draw_idle()

    @staticmethod
    def _zoom_at(ax, x, y, scale):
        if x is None or y is None:
            return
        x0, x1 = ax.get_xlim(); y0, y1 = ax.get_ylim()
        w = (x1 - x0) / scale; h = (y1 - y0) / scale
        rx = (x - x0) / (x1 - x0) if (x1 - x0) else 0.5
        ry = (y - y0) / (y1 - y0) if (y1 - y0) else 0.5
        ax.set_xlim(x - w * rx, x + w * (1 - rx))
        ax.set_ylim(y - h * ry, y + h * (1 - ry))

    # --- drag-to-pan (left mouse button) ---
    def _on_press(self, event):
        if event.button != 1 or event.inaxes != self.ax:
            return
        ax = self.ax
        self._pan_state = (event.xdata, event.ydata, ax.get_xlim(), ax.get_ylim())

    def _on_motion(self, event):
        if self._pan_state is None or event.inaxes != self.ax:
            return
        x0, y0, (xl0, xl1), (yl0, yl1) = self._pan_state
        if (event.xdata is None) or (event.ydata is None):
            return
        dx = event.xdata - x0
        dy = event.ydata - y0
        # pan by shifting original limits opposite to mouse move
        self.ax.set_xlim(xl0 - dx, xl1 - dx)
        self.ax.set_ylim(yl0 - dy, yl1 - dy)
        self.canvas.draw_idle()

    def _on_release(self, event):
        if event.button == 1:
            self._pan_state = None

    # ---------- teardown ----------
    def _on_close(self):
        # disconnect mpl events
        for cid in self._mpl_cids:
            try:
                self.canvas.mpl_disconnect(cid)
            except Exception:
                pass
        self._mpl_cids.clear()
        try:
            self.canvas.get_tk_widget().destroy()
        except Exception:
            pass
        try:
            import matplotlib.pyplot as plt
            plt.close(self.fig)
        except Exception:
            pass
        try:
            self.destroy()
        except Exception:
            pass



# ------------------------- Frames -------------------------
class BaseFrame(ttk.Frame):
    def __init__(self, parent, controller: App):
        super().__init__(parent)
        self.controller = controller

class WelcomeFrame(BaseFrame):
    def __init__(self, parent, controller):
        super().__init__(parent, controller)
        ttk.Label(self, text="Welcome", font=("Segoe UI", 18, "bold")).pack(anchor="w", padx=16, pady=(16, 8))
        ttk.Label(
            self,
            text=("This window runs the whole pipeline: choose a DOCX, reuse or create a concentration workbook, "
                  "assign concentrations, validate, compute equilibrium, and export results."),
            justify="left",
        ).pack(anchor="w", padx=16)

class DocxFrame(BaseFrame):
    def __init__(self, parent, controller):
        super().__init__(parent, controller)
        ttk.Label(self, text="Step 1 — Pick DOCX & Parse", font=("Segoe UI", 14, "bold")).pack(anchor="w", padx=16, pady=(16, 8))
        row = ttk.Frame(self); row.pack(fill="x", padx=16)
        ttk.Entry(row, textvariable=controller.state["docx_path"], width=80).pack(side="left", fill="x", expand=True)
        ttk.Button(row, text="Browse", command=self._pick).pack(side="left", padx=8)
        ttk.Button(self, text="Parse", command=self._parse).pack(anchor="e", padx=16, pady=8)
        self.msg = ttk.Label(self, text="")
        self.msg.pack(anchor="w", padx=16, pady=4)

    def _pick(self):
        path = filedialog.askopenfilename(
            title="Select reactions .docx",
            filetypes=[("Word document", "*.docx")],
            initialdir=str(BASE_DIR),
        )
        if path:
            self.controller.state["docx_path"].set(path)

    def _parse(self):
        ok = self.controller.parse_docx()
        self.msg.config(text="Parsed systems." if ok else "Parse failed.")
        if ok:
            self.controller.next_frame()

class DataframeFrame(BaseFrame):
    """
    Step 2 — scan for previous .xlsx OR create/edit a new one inline.

    Left:  folder + Scan + matches list (select = chosen_xlsx immediately)
    Right: inline editor (Create new) with editable tables per system
    """
    def __init__(self, parent, controller):
        super().__init__(parent, controller)
        ttk.Label(self, text="Step 2 — Dataframe Selection",
                  font=("Segoe UI", 14, "bold")).pack(anchor="w", padx=16, pady=(16, 8))

        # --- Top controls: folder + scan ---
        top = ttk.Frame(self); top.pack(fill="x", padx=16)
        ttk.Label(top, text="Folder:").pack(side="left")
        ttk.Entry(top, textvariable=controller.state["reactants_folder"], width=60)\
            .pack(side="left", fill="x", expand=True, padx=8)
        ttk.Button(top, text="Browse", command=self._pick_folder).pack(side="left")
        ttk.Button(top, text="Scan", command=self._scan).pack(side="left", padx=8)

        # Auto-scan whenever the folder StringVar changes
        controller.state["reactants_folder"].trace_add("write", lambda *a: self._scan())

        # --- Split main area: left (matches) | right (inline editor) ---
        main = ttk.Panedwindow(self, orient="horizontal")
        main.pack(fill="both", expand=True, padx=16, pady=8)

        # Left: matches list
        left = ttk.Frame(main)
        ttk.Label(left, text="Matching previous files").pack(anchor="w")
        self.listbox = tk.Listbox(left, height=16, exportselection=False)
        self.listbox.pack(fill="both", expand=True, pady=(4, 8))
        self.listbox.bind("<<ListboxSelect>>", self._select_file)
        self.listbox.bind("<Double-Button-1>", lambda e: self._goto_next())

        # Keyboard navigation
        for seq in ("<Up>", "<Left>"):
            self.listbox.bind(seq, lambda e: self._nudge(-1))
        for seq in ("<Down>", "<Right>"):
            self.listbox.bind(seq, lambda e: self._nudge(+1))
        self.listbox.bind("<Tab>", lambda e: self._nudge(+1))
        # Shift+Tab has two spellings depending on platform
        self.listbox.bind("<Shift-Tab>", lambda e: self._nudge(-1))
        self.listbox.bind("<ISO_Left_Tab>", lambda e: self._nudge(-1))
        self.listbox.bind("<Home>", lambda e: self._jump_first())
        self.listbox.bind("<End>", lambda e: self._jump_last())
        self.listbox.bind("<Return>", lambda e: self._goto_next())

        lrow = ttk.Frame(left); lrow.pack(fill="x")
        ttk.Button(lrow, text="Create new (inline)", command=self._start_inline_new).pack(side="left")
        ttk.Button(lrow, text="Clear selection", command=self._clear_choice).pack(side="left", padx=8)

        main.add(left, weight=1)

        # Right: inline editor with callback when saving
        self.editor_panel = EditorPanel(main, controller, on_saved=self._on_saved_newfile)
        main.add(self.editor_panel, weight=2)

        # Bottom status
        self.status = ttk.Label(self, text="Chosen file: <none>")
        self.status.pack(anchor="w", padx=16, pady=(0, 12))

    def on_show(self):
        self._scan()
        # focus list for immediate keyboard navigation
        self.after(50, lambda: self.listbox.focus_set())

    # ---------- left side behaviours ----------
    def _pick_folder(self):
        from tkinter.filedialog import askdirectory
        folder = askdirectory(title="Select folder for concentration dataframes")
        if folder:
            self.controller.state["reactants_folder"].set(folder)  # trace calls _scan()

    def _scan(self):
        if not self.controller.compartment_manager:
            self.listbox.delete(0, tk.END)
            self.listbox.insert(tk.END, "Parse a DOCX first (Step 1).")
            return

        matches = self.controller.detect_previous_frames()
        self.listbox.delete(0, tk.END)
        if matches:
            for m in matches:
                self.listbox.insert(tk.END, m)
            # If nothing chosen, preselect first match
            if not self.controller.state["chosen_xlsx"].get().strip():
                self.listbox.selection_set(0)
                self._select_file()
        else:
            self.listbox.insert(tk.END, "No matches found.")
        self._sync_status()

    # Keyboard helpers
    def _valid_indices(self):
        """Indices that are real files (skip placeholder lines)."""
        vals = [self.listbox.get(i) for i in range(self.listbox.size())]
        return [i for i, v in enumerate(vals)
                if v not in ("No matches found.", "Parse a DOCX first (Step 1).")]

    def _current_index(self):
        sel = self.listbox.curselection()
        return sel[0] if sel else None

    def _nudge(self, delta: int):
        """Move selection by +1 or -1 without wrap-around."""
        valid = self._valid_indices()
        if not valid:
            return "break"
        cur = self._current_index()
        # If nothing selected, start at first valid
        if cur is None or cur not in valid:
            target = valid[0] if delta >= 0 else valid[-1]
        else:
            pos = valid.index(cur)
            new_pos = pos + delta
            if new_pos < 0 or new_pos >= len(valid):
                return "break"  # linear ends; no wrap
            target = valid[new_pos]
        self.listbox.selection_clear(0, tk.END)
        self.listbox.selection_set(target)
        self.listbox.see(target)
        self._select_file()
        return "break"  # prevent default listbox behavior

    def _jump_first(self):
        valid = self._valid_indices()
        if not valid: return "break"
        idx = valid[0]
        self.listbox.selection_clear(0, tk.END)
        self.listbox.selection_set(idx)
        self.listbox.see(idx)
        self._select_file()
        return "break"

    def _jump_last(self):
        valid = self._valid_indices()
        if not valid: return "break"
        idx = valid[-1]
        self.listbox.selection_clear(0, tk.END)
        self.listbox.selection_set(idx)
        self.listbox.see(idx)
        self._select_file()
        return "break"

    def _select_file(self, event=None):
        sel = self.listbox.curselection()
        if not sel:
            return
        fname = self.listbox.get(sel[0])
        if fname in ("No matches found.", "Parse a DOCX first (Step 1)."):
            return
        folder = self.controller.state["reactants_folder"].get().strip()
        self.controller.state["chosen_xlsx"].set(os.path.join(folder, fname))
        # Hide editor if a previous file is chosen
        self.editor_panel.reset()
        self._sync_status()

    def _clear_choice(self):
        self.controller.state["chosen_xlsx"].set("")
        self._sync_status()

    def _goto_next(self):
        if self.controller.state["chosen_xlsx"].get().strip():
            self.controller.next_frame()

    # ---------- right side behaviours (inline editor) ----------
    def _start_inline_new(self):
        if not self.controller.compartment_manager:
            messagebox.showerror("Not ready", "Parse a DOCX first (Step 1).")
            return
        cm = self.controller.compartment_manager
        systems = cm.get_all_systems()
        dfs = {}
        for sys_name, info in systems.items():
            df = info["speciesMatrix"].construct_concentration_data(system_name=sys_name)
            if "concentration" not in df.columns:
                df["concentration"] = pd.NA
            dfs[sys_name] = df[["species", "concentration"]]
        self.editor_panel.load(dfs)
        # Clear any previously chosen file while editing inline
        self.controller.state["chosen_xlsx"].set("")
        self._sync_status()

    def _on_saved_newfile(self, path: str):
        """Called by the EditorPanel after saving a new workbook."""
        self.controller.state["chosen_xlsx"].set(path)
        self._scan()  # refresh matches
        fname = os.path.basename(path)
        for i in range(self.listbox.size()):
            if self.listbox.get(i) == fname:
                self.listbox.selection_clear(0, tk.END)
                self.listbox.selection_set(i)
                self.listbox.see(i)
                break
        self.status.config(text=f"Chosen file: {fname}")
        # bring focus back so arrow/tab continue working
        self.listbox.focus_set()

    # ---------- helpers ----------
    def _sync_status(self):
        chosen = self.controller.state["chosen_xlsx"].get().strip()
        self.status.config(text=f"Chosen file: {os.path.basename(chosen) if chosen else '<none>'}")

class EditorPanel(ttk.Frame):
    """
    Inline editor with a tab per system.
    Each tab shows a 2-column editable table (species, concentration).
    On save, writes an .xlsx (one sheet per system) and notifies parent via callback.
    """
    def __init__(self, parent, controller: App, on_saved=None):
        super().__init__(parent)
        self.controller = controller
        self.on_saved = on_saved

        # Notebook for system tabs
        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill="both", expand=True)
        self.tables = {}  # sys_name -> tree widget

        # Save/discard row
        row = ttk.Frame(self)
        row.pack(fill="x", pady=6)
        ttk.Button(row, text="Save as .xlsx", command=self._save).pack(side="right", padx=8)
        ttk.Button(row, text="Discard", command=self.reset).pack(side="right")

    def reset(self):
        """Clear all tabs and tables."""
        for tab_id in self.notebook.tabs():
            self.notebook.forget(tab_id)
        self.tables.clear()

    def load(self, df_map: dict):
        """Load a dict of system_name -> DataFrame(species, concentration)."""
        self.reset()
        for sys_name, df in df_map.items():
            frame = ttk.Frame(self.notebook)
            self.notebook.add(frame, text=sys_name)
            tree = self._make_table(frame, df)
            self.tables[sys_name] = tree

    def _make_table(self, parent, df: pd.DataFrame):

        cols = ("species", "concentration")
        tree = ttk.Treeview(parent, columns=cols, show="headings")

        for c in cols:
            tree.heading(c, text=c)
            tree.column(c, width=200 if c == "species" else 160, anchor="w")

        # Insert rows
        for _, row in df.iterrows():
            tree.insert("", "end", values=(row.get("species", ""), row.get("concentration", "")))
        tree.pack(fill="both", expand=True)

        # --- Excel-like inline editing for the concentration column ---
        editor = {"entry": None, "item": None}

        def _select_row(item_id):
            tree.selection_set(item_id)
            tree.focus(item_id)
            tree.see(item_id)

        def _row_siblings():
            return list(tree.get_children(""))

        def _row_index(item_id):
            try:
                return _row_siblings().index(item_id)
            except ValueError:
                return -1

        def _neighbor(item_id, delta):
            sibs = _row_siblings()
            idx = _row_index(item_id)
            nxt = idx + delta
            if 0 <= nxt < len(sibs):
                return sibs[nxt]
            return None

        def _place_editor(item_id):
            """Overlay an Entry on the concentration cell, with select-all."""
            # compute bbox for column #2 (concentration)
            bbox = tree.bbox(item_id, column="#2")
            if not bbox:
                tree.see(item_id)
                bbox = tree.bbox(item_id, column="#2")
                if not bbox:
                    return "break"
            x, y, w, h = bbox

            cur = tree.set(item_id, "concentration")

            # remove any existing editor
            if editor["entry"] is not None:
                try:
                    editor["entry"].destroy()
                except Exception:
                    pass

            e = ttk.Entry(tree)
            e.insert(0, "" if cur is None else str(cur))
            e.place(x=x, y=y, width=w, height=h)
            e.focus_set()

            # IMPORTANT: select all to mimic Excel
            try:
                e.selection_range(0, tk.END)
                e.icursor(tk.END)  # caret at end (selection still active)
            except Exception:
                pass

            editor["entry"] = e
            editor["item"] = item_id

            def commit_and_move(delta=None):
                """Commit the edit and optionally move to another row and reopen editor."""
                val = e.get()
                tree.set(item_id, "concentration", val)
                try:
                    e.destroy()
                except Exception:
                    pass
                editor["entry"] = None
                editor["item"] = None

                if delta is not None:
                    nxt = _neighbor(item_id, delta)
                    if nxt is not None:
                        _select_row(nxt)
                        _place_editor(nxt)

            def cancel(_=None):
                try:
                    e.destroy()
                except Exception:
                    pass
                editor["entry"] = None
                editor["item"] = None
                return "break"

            # Key bindings inside the editor (Excel-ish)
            # Enter: commit & go down; Shift+Enter: commit & go up
            e.bind("<Return>", lambda _=None: (commit_and_move(+1), "break"))
            e.bind("<KP_Enter>", lambda _=None: (commit_and_move(+1), "break"))
            e.bind("<Shift-Return>", lambda _=None: (commit_and_move(-1), "break"))

            # Tab moves down; Shift+Tab moves up
            e.bind("<Tab>", lambda _=None: (commit_and_move(+1), "break"))
            e.bind("<ISO_Left_Tab>", lambda _=None: (commit_and_move(-1), "break"))  # Linux Shift+Tab
            e.bind("<Shift-Tab>", lambda _=None: (commit_and_move(-1), "break"))

            # Arrow keys move to neighbor cells (since only one editable column)
            e.bind("<Down>", lambda _=None: (commit_and_move(+1), "break"))
            e.bind("<Up>",   lambda _=None: (commit_and_move(-1), "break"))
            e.bind("<Right>",lambda _=None: (commit_and_move(+1), "break"))
            e.bind("<Left>", lambda _=None: (commit_and_move(-1), "break"))

            # Esc cancels
            e.bind("<Escape>", cancel)

            return "break"

        # Mouse: single click selects row; Enter will start editing
        def on_click(event):
            item = tree.identify_row(event.y)
            if not item:
                return
            _select_row(item)
            # If clicked the concentration cell directly, you can choose to start edit on single click:
            # col = tree.identify_column(event.x)
            # if col == "#2":  # uncomment if you want single-click to enter edit
            #     return _place_editor(item)

        tree.bind("<Button-1>", on_click)

        # Double-click anywhere in the row → start editing concentration (Excel feel)
        def on_double_click(event):
            item = tree.identify_row(event.y)
            if not item:
                return
            _select_row(item)
            return _place_editor(item)

        tree.bind("<Double-1>", on_double_click)

        # Keyboard: Enter/F2 opens editor with select-all on focused row
        def start_edit_current(_=None):
            item = tree.focus()
            if not item:
                sibs = _row_siblings()
                if not sibs: 
                    return "break"
                item = sibs[0]
                _select_row(item)
            return _place_editor(item)

        tree.bind("<Return>", start_edit_current)
        tree.bind("<KP_Enter>", start_edit_current)
        tree.bind("<F2>", start_edit_current)

        # When NOT editing: arrow keys move selection and open editor at new row (select-all)
        def move_and_edit(delta):
            if editor["entry"] is not None:
                return "break"  # editing handled by entry bindings
            item = tree.focus()
            if not item:
                sibs = _row_siblings()
                if not sibs:
                    return "break"
                item = sibs[0]
            nxt = _neighbor(item, delta)
            if nxt is None:
                return "break"
            _select_row(nxt)
            return _place_editor(nxt)

        tree.bind("<Down>", lambda e: move_and_edit(+1))
        tree.bind("<Up>",   lambda e: move_and_edit(-1))
        tree.bind("<Right>",lambda e: move_and_edit(+1))
        tree.bind("<Left>", lambda e: move_and_edit(-1))

        # Tab behavior when not editing: move to next and start edit
        def tab_move(_=None):
            return move_and_edit(+1)
        def shifttab_move(_=None):
            return move_and_edit(-1)
        tree.bind("<Tab>", tab_move)
        tree.bind("<ISO_Left_Tab>", shifttab_move)
        tree.bind("<Shift-Tab>", shifttab_move)

        return tree



    def _save(self):
        """Export current tables to an Excel workbook and set as chosen_xlsx."""
        data_map = {}
        for sys_name, tree in self.tables.items():
            rows = []
            for iid in tree.get_children(""):
                species, conc = tree.item(iid, "values")
                rows.append({"species": species, "concentration": conc})
            data_map[sys_name] = pd.DataFrame(rows, columns=["species", "concentration"])

        if not data_map:
            messagebox.showwarning("Nothing to save", "No tables to save.")
            return

        folder = self.controller.state["reactants_folder"].get().strip() or "."
        initial = os.path.join(folder, "concentration_requesting_data.xlsx")
        out = filedialog.asksaveasfilename(
            title="Save concentration workbook as",
            defaultextension=".xlsx",
            initialfile=os.path.basename(initial),
            initialdir=os.path.dirname(initial),
            filetypes=[("Excel files", "*.xlsx")],
        )
        if not out:
            return

        with pd.ExcelWriter(out, engine="openpyxl") as writer:
            for sys_name, df in data_map.items():
                df.to_excel(writer, sheet_name=sys_name, index=False)

        # Mark this file as chosen and notify parent
        self.controller.state["chosen_xlsx"].set(out)
        messagebox.showinfo("Saved", f"Workbook saved to:\n{out}")
        if callable(self.on_saved):
            self.on_saved(out)

class RunFrame(BaseFrame):
    def __init__(self, parent, controller):
        super().__init__(parent, controller)
        ttk.Label(self, text="Step 3 — Assign, Validate & Run", font=("Segoe UI", 14, "bold")).pack(anchor="w", padx=16, pady=(16, 8))
        ttk.Button(self, text="Assign concentrations from workbook", command=self._assign).pack(anchor="w", padx=16, pady=6)
        ttk.Button(self, text="Run compute & export", command=self._run).pack(anchor="w", padx=16, pady=6)
        ttk.Button(self, text="Show convergence plots", command=self.controller.show_convergence_plots)\
            .pack(anchor="w", padx=16, pady=6)
        self.log = tk.Text(self, height=18, wrap="word")
        self.log.pack(fill="both", expand=True, padx=16, pady=12)

    def append(self, text: str):
        self.log.insert(tk.END, text + "\n")
        self.log.see(tk.END)

    def _assign(self):
        ok = self.controller.assign_concentrations()
        self.append("Assigned concentrations." if ok else "Assign failed.")

    def _run(self):
        try:
            ok = self.controller.run_compute_and_export()
            self.append("Compute finished and outputs exported." if ok else "Run failed.")
        except Exception as e:
            messagebox.showerror("Run error", str(e))
            self.append(f"ERROR: {e}")

# ------------------------- Entrypoint -------------------------
if __name__ == "__main__":
    App().mainloop()
